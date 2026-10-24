import json
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from site_bot import chat_export
from site_bot.classifier import Classification, ClaudeClassifier, KeywordClassifier

ANDROID_EN = """24/10/2026, 08:01 - Messages and calls are end-to-end encrypted. No one outside of this chat can read them.
24/10/2026, 08:02 - Ahmad added Samer
24/10/2026, 09:15 - م. سامر: بدنا 20 طن حديد قطر 12 لبكرا
موقع بلوك B
24/10/2026, 09:20 - Ahmad: تمام
24/10/2026, 09:21 - Ahmad: <Media omitted>
24/10/2026, 13:40 - Ahmad: الحديد انطلب من المورد وبكرا الصبح بوصل
25/10/2026, 07:55 - م. سامر: This message was deleted
"""

# Arabic-locale Android: right-to-left marks, Arabic-Indic digits, Arabic comma, ص/م
ANDROID_AR = (
    "‏٢٤‏/١٠‏/٢٠٢٦، ٢:٠٥ م - م. سامر: صبينا سقف الطابق التاني\n"
    "‏٢٤‏/١٠‏/٢٠٢٦، ١١:٣٠ ص - المشتريات: <تم استبعاد الوسائط>\n"
)

IOS = (
    "[10/24/26, 2:05:13 PM] Samer: ‎<attached: 00000012-PHOTO-2026-10-24-14-05-13.jpg>\n"
    "[10/24/26, 2:06:01 PM] Samer: الباطون ما وصل\n"
    "[10/25/26, 9:00:00 AM] Ahmad: ok\n"
)


def test_android_english_export():
    msgs = chat_export.parse(ANDROID_EN)
    assert [m.sender for m in msgs] == ["م. سامر", "Ahmad", "Ahmad", "Ahmad", "م. سامر"]
    assert msgs[0].text == "بدنا 20 طن حديد قطر 12 لبكرا\nموقع بلوك B"
    assert msgs[0].sent_at == datetime(2026, 10, 24, 9, 15)
    assert msgs[2].media and msgs[4].deleted
    assert [m.needs_ai for m in msgs] == [True, False, False, True, False]


def test_android_arabic_locale_export():
    msgs = chat_export.parse(ANDROID_AR)
    assert msgs[0].sent_at == datetime(2026, 10, 24, 14, 5)
    assert msgs[0].text == "صبينا سقف الطابق التاني"
    assert msgs[1].sent_at == datetime(2026, 10, 24, 11, 30) and msgs[1].media


def test_ios_export_month_first():
    msgs = chat_export.parse(IOS)
    assert msgs[0].media and msgs[0].sent_at == datetime(2026, 10, 24, 14, 5)
    assert msgs[1].text == "الباطون ما وصل"
    assert not msgs[2].needs_ai


def test_classify_only_sends_real_messages_with_context():
    msgs = chat_export.parse(ANDROID_EN)
    seen = []

    class Spy:
        def classify_batch(self, messages, earlier, projects):
            seen.append(([m["index"] for m in messages], [m["index"] for m in earlier]))
            return [Classification("other", 0.9, m["text"]) for m in messages]

    results = chat_export.classify(msgs, Spy(), [], batch_size=1, context=3)
    assert seen == [([0], []), ([3], [0, 1])]  # media (index 2) is left out of the context
    assert [r is None for r in results] == [False, True, True, False, True]


def test_review_workbook(tmp_path):
    msgs = chat_export.parse(ANDROID_EN)
    results = chat_export.classify(msgs, KeywordClassifier(), [{"code": "P03", "name": "عمارة"}])
    out = tmp_path / "review.xlsx"
    chat_export.write_review(out, msgs, results)

    wb = load_workbook(out)
    assert wb.sheetnames == ["ملخص", "الرسائل", "طلبات المواد"]
    ws = wb["الرسائل"]
    assert ws.sheet_view.rightToLeft and ws.max_row == 6
    assert ws["E2"].value == "طلب مواد" and ws["E3"].value == "دردشة (تخطّي)"
    assert ws["E5"].value == "تحديث مشتريات"
    req = wb["طلبات المواد"]
    assert [c.value for c in req[2]][3:7] == ["حديد", "Ø12", 20, "طن"]


def test_claude_batch_keeps_order_and_flags_missing():
    def result(i, cat):
        return {"index": i, "category": cat, "confidence": 0.9, "summary": "s", "project_code": None,
                "block": None, "floor": None, "items": [], "needed_by": None, "boq_code": None,
                "workers": None, "equipment": None}

    def create(**kw):
        assert "<earlier>" in kw["messages"][0]["content"]
        text = json.dumps({"results": [result(7, "issue"), result(5, "progress")]})
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=text)])

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    msgs = [{"index": i, "sender": "x", "time": "t", "text": "m"} for i in (5, 6, 7)]
    got = ClaudeClassifier(client=client).classify_batch(msgs, [], [])
    assert [c.category for c in got] == ["progress", "other", "issue"]
    assert got[1].confidence == 0.0
