"""Read a WhatsApp "Export chat" .txt file, classify its messages, and write an Excel sheet to review.

The review sheet is how classification accuracy gets measured on real group chats before anything
is built on top of them: someone marks each row right or wrong, and the summary sheet counts.
"""
import re
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .classifier import CATEGORY_AR, Classification

_INVISIBLE = re.compile("[‎‏‪-‮⁦-⁩﻿]")
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_STAMP = (r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4}),?\s+(\d{1,2}):(\d{2})(?::\d{2})?"
          r"\s*([AaPp]\.?\s?[Mm]\.?|ص|م)?")
_ANDROID = re.compile(rf"^{_STAMP}\s+-\s+(.*)$")       # 24/10/2026, 14:05 - Name: text
_IOS = re.compile(rf"^\[{_STAMP}\]\s+(.*)$")            # [24/10/2026, 2:05:13 PM] Name: text
# A media marker opens the message; whatever follows it is the caption (or just the file name).
_MEDIA = re.compile(
    r"^(?:\[Forwarded\]\s*)?(?:<[^<>\n]*omitted>|<album message>|<attached: [^>]*>|<تم استبعاد [^<>\n]*>|"
    r"(?:image|video|audio|sticker|GIF|document) omitted|\S+ \(file attached\)|الوسائط غير مضمّنة)",
    re.I,
)
_FILENAME = re.compile(r"^\S.*\.(pdf|xlsx?|docx?|pptx?|jpe?g|png|heic|mp4|opus|m4a)$", re.I)
_CONTACT = re.compile(r"^(?:\[Forwarded\]\s*)?Contact: BEGIN:VCARD|^BEGIN:VCARD")
_DELETED = re.compile(r"^(This message was deleted|You deleted this message|تم حذف هذه الرسالة|لقد حذفت هذه الرسالة)\.?$")
_CHATTER = re.compile(
    r"^(?:(?:ok|okay|noted|thanks|thank you|yes|sure|aameen|ameen|alhamdulillah|alhumdulillah|"
    r"تمام|تم|اوكي|اوك|شكرا|شكراً|مشكور|يعطيك العافيه|يعطيك العافية|الله يعطيك العافيه|اللهم امين|امين)"
    r"(?:[ ,.]+(?:sir|madam|noted|thanks|engineer))*(?:[ ,.]+\w+)?|<album message>|(?:@\S+[ ,]*)+)[.!\s👍🙏✅🌹🤲🏻]*$",
    re.I,
)


@dataclass
class ChatMessage:
    sent_at: datetime
    sender: str
    text: str
    media: bool = False
    deleted: bool = False
    contact: bool = False

    @property
    def caption(self) -> str:
        """The words someone typed: the text without a leading media marker or bare file name."""
        body = _MEDIA.sub("", self.text, count=1).strip() if self.media else self.text
        return "" if _FILENAME.match(body) else body

    @property
    def needs_ai(self) -> bool:
        body = self.caption
        return not (self.deleted or self.contact or len(body) < 3 or _CHATTER.match(body))


def parse(raw: str) -> list[ChatMessage]:
    lines = _INVISIBLE.sub("", raw.translate(_DIGITS)).replace(" ", " ").replace("،", ",").splitlines()
    stamped = []  # (date parts, rest) for every line that starts a message
    for line in lines:
        m = _ANDROID.match(line) or _IOS.match(line)
        if m:
            stamped.append([m.groups(), [m.group(7)]])
        elif stamped:
            stamped[-1][1].append(line)  # continuation of a multi-line message

    day_first = _day_first([g for g, _ in stamped])
    out = []
    for g, rest_lines in stamped:
        head, sep, first = rest_lines[0].partition(": ")
        if not sep:
            continue  # system line: "X added Y", encryption notice, ...
        text = "\n".join([first, *rest_lines[1:]]).strip()
        out.append(ChatMessage(_when(g, day_first), head.strip(), text, media=bool(_MEDIA.match(text)),
                               deleted=bool(_DELETED.match(text)), contact=bool(_CONTACT.match(text))))
    return out


def _day_first(groups) -> bool:
    if any(int(g[1]) > 12 for g in groups):
        return False
    return True  # also the default when ambiguous: day-first is the norm in the region


def _when(g, day_first: bool) -> datetime:
    a, b, y, hh, mm, ampm = g[:6]
    day, month = (int(a), int(b)) if day_first else (int(b), int(a))
    year = int(y) + 2000 if len(y) == 2 else int(y)
    hour = int(hh)
    if ampm:
        pm = ampm.lower().startswith("p") or ampm == "م"
        hour = hour % 12 + (12 if pm else 0)
    return datetime(year, month, day, hour, int(mm))


def classify(messages: list[ChatMessage], classifier, projects: list[dict],
             batch_size: int = 30, context: int = 8, progress=None) -> list[Classification | None]:
    """One result per message; None for media, deleted and chatter (not sent to the AI)."""
    results: list[Classification | None] = [None] * len(messages)
    todo = [i for i, m in enumerate(messages) if m.needs_ai]
    as_item = lambda i: {"index": i, "sender": messages[i].sender,
                         "time": messages[i].sent_at.strftime("%Y-%m-%d %H:%M"), "text": messages[i].text}
    for start in range(0, len(todo), batch_size):
        chunk = todo[start:start + batch_size]
        first = chunk[0]
        earlier = [as_item(i) for i in range(max(0, first - context), first)
                   if not messages[i].contact and (messages[i].caption or not messages[i].media)]
        for i, c in zip(chunk, classifier.classify_batch([as_item(i) for i in chunk], earlier, projects)):
            results[i] = c
        if progress:
            progress(min(start + batch_size, len(todo)), len(todo))
    return results


def batches_needed(messages: list[ChatMessage], batch_size: int = 30) -> int:
    n = sum(m.needs_ai for m in messages)
    return -(-n // batch_size)


def write_review(path, messages: list[ChatMessage], results: list[Classification | None]) -> None:
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5566")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, title, headers, widths):
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for col, w in zip("ABCDEFGHIJKLMNOP", widths):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    sheet(ws, "الرسائل",
          ["التاريخ", "الوقت", "المرسل", "الرسالة", "التصنيف", "الثقة", "المشروع", "الملخص",
           "المواد", "الموعد", "التصنيف صح؟", "التصنيف الصحيح", "ملاحظات"],
          [11, 7, 16, 50, 15, 7, 9, 36, 30, 12, 12, 16, 24])
    reviewed = 0
    for m, c in zip(messages, results):
        if c is None:
            label = ("جهة اتصال" if m.contact else "محذوفة" if m.deleted
                     else "وسائط" if m.media and not m.caption else "دردشة (تخطّي)")
            row = [label, None, None, None, None, None]
        else:
            reviewed += 1
            items = "، ".join(" ".join(str(x) for x in (it["material"], it["spec"], it["quantity"] and f"{it['quantity']:g}",
                                                       it["unit"]) if x) for it in c.items)
            row = [CATEGORY_AR.get(c.category, c.category), round(c.confidence, 2), c.project_code,
                   c.summary, items or None, c.needed_by]
        text = m.text.split("\n", 1)[0][:60] + " …" if m.contact else m.text
        ws.append([m.sent_at.date(), m.sent_at.strftime("%H:%M"), m.sender, text, *row])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = wrap
        r[0].number_format = "yyyy-mm-dd"
    last = ws.max_row
    yes_no = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
    cats = DataValidation(type="list", formula1='"' + ",".join(CATEGORY_AR.values()) + '"', allow_blank=True)
    ws.add_data_validation(yes_no); ws.add_data_validation(cats)
    yes_no.add(f"K2:K{last}"); cats.add(f"L2:L{last}")

    req = wb.create_sheet()
    sheet(req, "طلبات المواد", ["التاريخ", "المرسل", "المشروع", "المادة", "المواصفة", "الكمية", "الوحدة",
                                "الموعد", "الرسالة"], [11, 16, 9, 18, 12, 9, 8, 12, 50])
    for m, c in zip(messages, results):
        if c and c.category == "material_request":
            for it in c.items or [{"material": c.summary, "spec": None, "quantity": None, "unit": None}]:
                req.append([m.sent_at.date(), m.sender, c.project_code, it["material"], it["spec"],
                            it["quantity"], it["unit"], c.needed_by, m.text])
    for r in req.iter_rows(min_row=2):
        r[0].number_format = "yyyy-mm-dd"
        r[-1].alignment = wrap

    s = wb.create_sheet()
    sheet(s, "ملخص", ["البند", "القيمة"], [34, 40])
    first, lastd = (messages[0].sent_at.date(), messages[-1].sent_at.date()) if messages else (None, None)
    s.append(["الفترة", f"{first} إلى {lastd}"])
    s.append(["عدد الرسائل", len(messages)])
    s.append(["رسائل صنّفها النظام", reviewed])
    for key, name in CATEGORY_AR.items():
        s.append([f"  {name}", sum(1 for c in results if c and c.category == key)])
    s.append(["تمت مراجعتها", "=COUNTIF('الرسائل'!K:K,\"نعم\")+COUNTIF('الرسائل'!K:K,\"لا\")"])
    s.append(["دقة التصنيف", "=IFERROR(COUNTIF('الرسائل'!K:K,\"نعم\")/B" + str(s.max_row) + ",\"-\")"])
    s.cell(s.max_row, 2).number_format = "0%"
    s.append([])
    s.append(["طريقة المراجعة", "بعمود «التصنيف صح؟» اختار نعم أو لا لكل رسالة صنّفها النظام. "
                                "إذا لا، اختار التصنيف الصحيح. الدقة بتنحسب هون تلقائياً."])
    s.cell(s.max_row, 2).alignment = wrap
    wb.move_sheet("ملخص", offset=-2)
    wb.active = 0
    wb.save(path)
